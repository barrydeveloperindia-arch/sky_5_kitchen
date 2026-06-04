import json
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

# Load combos from JSON
combos_json_path = '03_Automation_System/combos.json'
with open(combos_json_path, 'r', encoding='utf-8') as f:
    combos = json.load(f)

# Load Excel workbook
excel_path = 'Hotel_Sky5_Market_Pricing_Audit.xlsx'
wb = openpyxl.load_workbook(excel_path)
sheet = wb.active

# Category mapping
cat_map = {
    'Breakfast': 'Breakfast',
    'Snacks': 'Snacks',
    'Salad': 'Salads',
    'Chinese': 'Chinese',
    'Rice': 'Rice',
    'Breads': 'Breads',
    'Thalis': 'Combos & Thalis',
    'Sweet Dish': 'Desserts',
    'Beverages': 'Beverages'
}

# Separate Veg/Non-Veg Main Course based on description icons
def get_main_course_category(item):
    if item['category'] == 'Main Course':
        if '🔴' in item['description']:
            return 'Non-Veg Main Course'
        else:
            return 'Veg Main Course'
    return cat_map.get(item['category'], item['category'])

# Create a map of items from combos.json
combos_map = {}
for item in combos:
    name_clean = item['name'].strip().lower()
    combos_map[name_clean] = item

print("Updating existing rows in Excel...")
updated_names = set()

# Iterate over existing rows (row 2 to the end)
max_row = sheet.max_row
for r in range(2, max_row + 1):
    item_name = sheet.cell(row=r, column=2).value
    if not item_name:
        continue
        
    item_name_clean = item_name.strip().lower()
    
    # Handle the old 'Butter Chicken' row rename to 'Butter Chicken (Half)'
    if item_name_clean == 'butter chicken':
        sheet.cell(row=r, column=2).value = 'Butter Chicken (Half)'
        item_name_clean = 'butter chicken (half)'
        
    # Find item in combos
    if item_name_clean in combos_map:
        item = combos_map[item_name_clean]
        sky5_price = item['price']
        market_price = sky5_price + 5
        diff = 5
        savings = (5 / market_price) * 100
        
        sheet.cell(row=r, column=3).value = sky5_price
        sheet.cell(row=r, column=4).value = market_price
        sheet.cell(row=r, column=5).value = diff
        sheet.cell(row=r, column=6).value = savings
        
        updated_names.add(item_name_clean)
        print(f"Row {r} updated: '{sheet.cell(row=r, column=2).value}' -> Sky5: {sky5_price}, Market: {market_price}")

# Determine items that need to be appended
to_append = []
for item in combos:
    name_clean = item['name'].strip().lower()
    if name_clean not in updated_names:
        to_append.append(item)

# Styles from the last data row to copy formatting
last_row = sheet.max_row
ref_cells = [sheet.cell(row=last_row, column=c) for c in range(1, 8)]

def copy_style(src_cell, dest_cell):
    dest_cell.font = Font(name=src_cell.font.name, size=src_cell.font.size, bold=src_cell.font.bold, italic=src_cell.font.italic, color=src_cell.font.color)
    dest_cell.alignment = Alignment(horizontal=src_cell.alignment.horizontal, vertical=src_cell.alignment.vertical, wrap_text=src_cell.alignment.wrap_text)
    dest_cell.border = Border(left=src_cell.border.left, right=src_cell.border.right, top=src_cell.border.top, bottom=src_cell.border.bottom)
    if src_cell.fill and src_cell.fill.fill_type:
        dest_cell.fill = PatternFill(fill_type=src_cell.fill.fill_type, start_color=src_cell.fill.start_color, end_color=src_cell.fill.end_color)

# Append new rows
next_row = last_row + 1
for item in to_append:
    cat = get_main_course_category(item)
    name = item['name']
    sky5_price = item['price']
    market_price = sky5_price + 5
    diff = 5
    savings = (5 / market_price) * 100
    
    # Decide Strategy
    strategy = 'Competitive Entry'
    if cat == 'Beverages':
        strategy = 'High Frequency Item'
    elif cat == 'Desserts':
        strategy = 'Sweet Upsell'
    elif 'Thali' in name or 'Combo' in name:
        strategy = 'High Value Combo'

    row_data = [cat, name, sky5_price, market_price, diff, savings, strategy]
    
    for c, val in enumerate(row_data, 1):
        cell = sheet.cell(row=next_row, column=c, value=val)
        # Copy style from reference row
        copy_style(ref_cells[c - 1], cell)
        
        # Apply standard currency formatting for prices and percentage for savings
        if c in [3, 4, 5]:
            cell.number_format = '₹#,##0'
        elif c == 6:
            cell.number_format = '0.00"%"'
            
    print(f"Row {next_row} appended: '{name}' in '{cat}' -> Sky5: {sky5_price}, Market: {market_price}")
    next_row += 1

# Save the workbook
wb.save(excel_path)
print("Excel sheet updated and saved successfully!")
